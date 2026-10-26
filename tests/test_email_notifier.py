import sys
import unittest
from datetime import date, datetime

from openpyxl import load_workbook

import database
import email_notifier
from utils import week_bounds

try:
    from .base import ServiceTrackerTestCase
except ImportError:  # `unittest discover -s tests` puts tests/ directly on sys.path
    from base import ServiceTrackerTestCase


class EmailNotifierTests(ServiceTrackerTestCase):
    def test_build_summary_body_and_attachment(self) -> None:
        reference = date(2026, 10, 26)  # Monday
        start, end = week_bounds(reference)
        database.add_customer({
            "customer_name": "Due This Week",
            "customer_location": "St Albans",
            "service_type": "Maintenance",
            "next_service_date": (start + date.resolution).isoformat(),
            "fault_status": "Active",
            "outstanding_invoices": 125.5,
        })
        database.add_customer({
            "customer_name": "Overdue Site",
            "customer_location": "Harpenden",
            "service_type": "Repair",
            "next_service_date": (reference - date.resolution * 10).isoformat(),
            "outstanding_invoices": 75,
        })
        summary = email_notifier.build_weekly_summary(reference)
        self.assertIn("Weekly Service Tracker Summary", summary.subject)
        self.assertIn("Due This Week", summary.body)
        self.assertIn("Overdue Site", summary.body)
        self.assertIn("Active faults", summary.body)
        self.assertIn("£200.50", summary.body)
        self.assertEqual(summary.start_date, start)
        self.assertEqual(summary.end_date, end)
        self.assertTrue(summary.attachment_path.is_file())

        workbook = load_workbook(summary.attachment_path)
        try:
            self.assertEqual(set(workbook.sheetnames), {
                "Summary", "Services Due", "Overdue Services", "Active Faults", "Outstanding Invoices",
            })
            self.assertEqual(workbook["Services Due"]["A2"].value, "Due This Week")
            self.assertEqual(workbook["Overdue Services"]["A2"].value, "Overdue Site")
        finally:
            workbook.close()

    def test_scheduled_email_predicate_is_once_per_day(self) -> None:
        now = datetime(2026, 10, 26, 9, 5)  # Monday, just after the target time
        settings = {
            "email_schedule_enabled": "1",
            "email_schedule_day": "monday",
            "email_schedule_time": "09:00",
            "email_last_sent": "",
        }
        self.assertTrue(email_notifier.should_send_scheduled(settings, now))
        settings["email_last_sent"] = "2026-10-26"
        self.assertFalse(email_notifier.should_send_scheduled(settings, now))
        settings["email_last_sent"] = ""
        self.assertFalse(email_notifier.should_send_scheduled(settings, datetime(2026, 10, 26, 8, 59)))
        self.assertFalse(email_notifier.should_send_scheduled(settings, datetime(2026, 10, 27, 9, 5)))
        email_notifier.mark_email_sent(now)
        self.assertEqual(database.get_setting("email_last_sent"), "2026-10-26")

    @unittest.skipIf(sys.platform == "win32", "Outlook may be installed on Windows test runners")
    def test_windows_only_outlook_send_reports_clear_error(self) -> None:
        self.assertFalse(email_notifier.is_outlook_available())
        summary = email_notifier.build_weekly_summary(date(2026, 10, 26))
        with self.assertRaises(email_notifier.EmailNotAvailableError):
            email_notifier.send_weekly_summary(summary, "manager@example.com")

    def test_recipient_validation_and_saved_settings(self) -> None:
        email_notifier.save_email_settings({
            "email_recipient": "manager@example.com",
            "email_schedule_enabled": True,
            "email_schedule_day": "friday",
            "email_schedule_time": "16:30",
            "not_allowed": "ignored",
        })
        settings = email_notifier.get_email_settings()
        self.assertEqual(settings["email_recipient"], "manager@example.com")
        self.assertEqual(settings["email_schedule_enabled"], "1")
        self.assertEqual(settings["email_schedule_day"], "friday")
        with self.assertRaises(ValueError):
            email_notifier.send_weekly_summary(
                email_notifier.build_weekly_summary(date(2026, 10, 26)), "invalid"
            )
