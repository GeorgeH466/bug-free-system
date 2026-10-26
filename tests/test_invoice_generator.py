from datetime import date
from pathlib import Path

from docx import Document

import database
import invoice_generator

try:
    from .base import ServiceTrackerTestCase
except ImportError:  # `unittest discover -s tests` puts tests/ directly on sys.path
    from base import ServiceTrackerTestCase

BUNDLED_TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "invoice_template.docx"


class InvoiceTests(ServiceTrackerTestCase):
    def test_bundled_template_has_expected_placeholders(self) -> None:
        self.assertTrue(BUNDLED_TEMPLATE.is_file())
        document = Document(str(BUNDLED_TEMPLATE))
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        for placeholder in (
            "{{customer_name}}", "{{customer_location}}", "{{id_code}}", "{{amount}}",
            "{{invoice_date}}", "{{service_type}}", "{{due_date}}",
        ):
            self.assertIn(placeholder, text)

    def test_generate_editable_invoice_and_record_payment(self) -> None:
        template = invoice_generator.ensure_default_template()
        self.assertTrue(template.is_file())
        customer_id = database.add_customer({
            "customer_name": "Acme Fire Systems",
            "customer_location": "12 High Street, St Albans",
            "id_code": "ACME-001",
            "service_type": "Maintenance",
            "outstanding_invoices": 175,
        })
        data = invoice_generator.default_invoice_data(
            database.get_customer(customer_id), date(2026, 10, 26)
        )
        data["line_items"] = [
            {"description": "Annual service", "quantity": 2, "unit_price": 75},
            {"description": "Replacement battery", "quantity": 1, "unit_price": 25},
        ]
        data["tax_rate"] = 20
        destination = self.workspace / "exports" / "Invoices" / "invoice.docx"
        result = invoice_generator.generate_invoice(
            data,
            destination,
            template_path=BUNDLED_TEMPLATE,
            update_payment=True,
        )
        self.assertTrue(destination.is_file())
        self.assertEqual(result["amount"], 210.0)
        self.assertEqual(len(database.get_payments_for_customer(customer_id)), 1)

        document = Document(str(destination))
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("Acme Fire Systems", text)
        self.assertIn("12 High Street, St Albans", text)
        self.assertIn("Annual service", text)
        self.assertIn("Replacement battery", text)
        self.assertIn("£210.00", text)
        self.assertEqual(len(document.tables[1].rows), 3)

    def test_invalid_invoice_rejected_before_file_write(self) -> None:
        destination = self.workspace / "invalid.docx"
        with self.assertRaises(ValueError):
            invoice_generator.generate_invoice({"customer_name": "No lines", "line_items": []}, destination)
        self.assertFalse(destination.exists())
