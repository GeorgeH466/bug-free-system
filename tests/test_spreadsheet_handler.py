from datetime import date

from openpyxl import Workbook, load_workbook

import database
from spreadsheet_handler import (
    ExportFilters,
    create_import_template,
    export_customers,
    import_customers,
)

try:
    from .base import ServiceTrackerTestCase
except ImportError:  # `unittest discover -s tests` puts tests/ directly on sys.path
    from base import ServiceTrackerTestCase


class SpreadsheetTests(ServiceTrackerTestCase):
    def _seed_customers(self) -> None:
        database.add_customer({
            "customer_name": "Northgate Alarms",
            "customer_location": "St Albans",
            "customer_email": "hello@northgate.example",
            "next_service_date": "2026-11-10",
            "fault_status": "Active",
            "outstanding_invoices": 125,
        })
        database.add_customer({
            "customer_name": "Clear Accounts",
            "customer_location": "Harpenden",
            "outstanding_invoices": 0,
        })

    def test_styled_export_with_filters(self) -> None:
        self._seed_customers()
        destination = self.workspace / "exports" / "filtered.xlsx"
        result = export_customers(destination, filters=ExportFilters(outstanding_only=True))
        self.assertEqual(result["row_count"], 1)
        workbook = load_workbook(destination)
        try:
            sheet = workbook["Customers"]
            self.assertEqual(sheet["B1"].value, "Customer Name")
            self.assertEqual(sheet["B2"].value, "Northgate Alarms")
            self.assertIsInstance(sheet["H2"].value, (date, type(None)))
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
            self.assertIn("Export Info", workbook.sheetnames)
        finally:
            workbook.close()

    def test_import_reports_valid_and_invalid_rows(self) -> None:
        source = self.workspace / "import.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Customer Name", "Customer Location", "Customer Email", "Customer Phone",
            "Survey Date", "Installation Date", "Last Service Date", "Next Service Date",
            "Service Type", "Outstanding Invoices",
        ])
        sheet.append([
            "Imported Client", "12 High Street", "valid@example.com", "01727 123456",
            "2025-01-01", "2025-01-10", "2026-01-10", "2027-01-10", "Maintenance", 0,
        ])
        sheet.append(["Bad Email", "Nowhere", "not-an-email", "ABC", "", "", "", "", "Repair", -1])
        sheet.append([
            "Bad Date Chain", "Nowhere", "okay@example.com", "01727 111111",
            "2025-05-01", "2025-04-01", "", "", "Repair", 0,
        ])
        workbook.save(source)
        workbook.close()

        dry = import_customers(source, dry_run=True)
        self.assertEqual(dry.imported, 1)
        self.assertEqual(dry.skipped, 2)
        self.assertIn("would be imported", dry.message)
        self.assertEqual(database.count_customers(), 0)

        summary = import_customers(source)
        self.assertEqual(summary.imported, 1)
        self.assertEqual(summary.skipped, 2)
        self.assertIn("Row 3", summary.details)
        self.assertEqual(database.count_customers(), 1)
        self.assertEqual(database.search_customers("Imported")[0]["next_service_date"], "2027-01-10")

    def test_import_real_excel_dates_and_alias_headers(self) -> None:
        source = self.workspace / "aliases.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Name", "Email", "Survey", "Installation", "Outstanding"])
        sheet.append(["Alias Client", "alias@example.com", date(2025, 1, 1), date(2025, 2, 1), 42])
        workbook.save(source)
        workbook.close()
        summary = import_customers(source)
        self.assertEqual(summary.imported, 1)
        customer = database.search_customers("Alias")[0]
        self.assertEqual(customer["survey_date"], "2025-01-01")
        self.assertEqual(customer["outstanding_invoices"], 42.0)

    def test_import_template_is_readable(self) -> None:
        path = create_import_template(self.workspace / "exports" / "template.xlsx")
        workbook = load_workbook(path)
        try:
            self.assertEqual(workbook["Customers"]["A1"].value, "Customer Name")
            self.assertIn("Instructions", workbook.sheetnames)
        finally:
            workbook.close()
